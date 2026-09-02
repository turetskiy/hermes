"""Low-level content writing into template paragraphs + the meta sidecar. fill_template.py owns the
STRUCTURE (finding/removing placeholders, roles, sections); this writes TEXT with the right fonts/accent
read from data/template.meta.json so filled content matches the chosen style. load_meta() runs first."""
import json
import os
import re
import docx
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

import paths
URL_RE = re.compile(r"(https?://\S+)")
META_PATH = os.path.join(paths.DATA, "template.meta.json")  # per-style sidecar from build_template

# Defaults match a sans template; load_meta() overrides them from the meta sidecar.
META = {}
BODY_FONT = "Calibri Light"
BOLD_FONT = "Calibri"
LINK_FONT = "Calibri Light"


def load_meta():
    global META, BODY_FONT, BOLD_FONT, LINK_FONT
    try:
        with open(META_PATH) as f:
            META = json.load(f)
    except FileNotFoundError:
        META = {}
    BODY_FONT = META.get("body_font", BODY_FONT)
    BOLD_FONT = META.get("bold_font", BOLD_FONT)
    LINK_FONT = META.get("link_font", LINK_FONT)


def dedash(text):
    return text.replace("—", "-").replace("–", "-")


def _char_spacing(run, twips):
    rPr = run._r.get_or_add_rPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(int(twips)))
    rPr.append(sp)


def add_hyperlink_run(p, url, text):
    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")            # keep Hyperlink style (blue+underline); override font+size only
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "Hyperlink")
    rPr.append(rstyle)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), LINK_FONT)
    rFonts.set(qn("w:hAnsi"), LINK_FONT)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")  # 9pt (links only)
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    hyperlink.append(r)
    p._p.append(hyperlink)


def clear_runs(p):
    el = p._element
    for child in list(el):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            el.remove(child)


def parse_bold_markup(text):
    segments = []
    for i, part in enumerate(text.split("**")):
        if part:
            segments.append((part, i % 2 == 1))
    return segments


def set_segments(p, text):
    clear_runs(p)
    for seg_text, bold in parse_bold_markup(dedash(text)):
        run = p.add_run(seg_text)
        run.bold = bold
        run.font.name = BOLD_FONT if bold else BODY_FONT


def apply_tagline_style(p):
    """Re-style the tagline runs to the template's tagline treatment (font/size/caps/spacing/colour).
    set_segments wrote the text in the body font; this restores the style's tagline design."""
    style = META.get("tagline")
    for r in p.runs:
        if style:
            if style.get("font"):
                r.font.name = style["font"]
            r.font.size = Pt(style.get("size", 10))
            if style.get("caps"):
                r.font.all_caps = True
            if style.get("color"):
                r.font.color.rgb = RGBColor.from_string(style["color"])
            if style.get("spacing") is not None:
                _char_spacing(r, style["spacing"])
        else:
            r.font.size = Pt(10)


def set_list_item(p, text):
    """Like set_segments but turns any URL into a real clickable hyperlink."""
    clear_runs(p)
    text = dedash(text)
    for part in URL_RE.split(text):
        if not part:
            continue
        if URL_RE.fullmatch(part):
            add_hyperlink_run(p, part, part)
        else:
            run = p.add_run(part)
            run.font.name = BODY_FONT
            if META.get("list_color"):
                run.font.color.rgb = RGBColor.from_string(META["list_color"])


